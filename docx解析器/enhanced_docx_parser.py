#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
增强版Docx解析器 - 解决样式继承问题
"""

import os
import tempfile
import requests
import json
import zipfile
import base64
import xml.etree.ElementTree as ET
from dataclasses import dataclass, asdict
from typing import Any, Dict, List, Optional
from urllib.parse import urlparse

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
except ImportError:
    print("错误: 需要安装python-docx库")
    print("请运行: pip install python-docx")
    exit(1)


@dataclass
class TextStyle:
    """字符样式"""
    font: str = None
    size: float = None
    color: str = None
    bold: bool = False
    italic: bool = False
    underline: bool = False
    subscript: bool = False
    superscript: bool = False
    highlight: str = None  # 突出显示颜色名称 (如 'yellow', 'cyan' 等)


@dataclass
class ParagraphStyle:
    """段落样式"""
    alignment: str = None
    line_spacing: float = None
    space_before: float = None
    space_after: float = None
    left_indent: float = None
    right_indent: float = None
    first_line_indent: float = None


@dataclass
class TableData:
    """表格数据"""
    rows: int
    columns: int
    cells: List[List[Optional[Dict[str, Any]]]]  # 修改为包含单元格详细信息的列表


@dataclass
class ImageData:
    """图片数据"""
    width: float
    height: float
    inline: bool
    alignment: str = None  # 图片对齐方式
    alt_text: str = None
    image_path: str = None  # 临时存储路径
    binary_content: str = None  # Base64编码的内容
    # 此类用于存储文档中提取的图片及其元数据


@dataclass
class DocumentElement:
    """文档元素"""
    element_type: str  # 'paragraph', 'table', 'image'
    content: Any
    paragraph_style: Optional[ParagraphStyle] = None
    text_style: Optional[TextStyle] = None


class EnhancedDocxParser:
    """增强版Docx文件解析器，解决样式继承问题"""
    
    def __init__(self):
        self.temp_dir = tempfile.mkdtemp()
        self.cache_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'image_cache')
        if not os.path.exists(self.cache_dir):
            os.makedirs(self.cache_dir)
        self.style_map = {}  # 样式ID到样式信息的映射
        self.style_cache = {}  # 样式计算缓存
    
    def cleanup(self):
        """清理临时目录"""
        import shutil
        if os.path.exists(self.temp_dir):
            try:
                shutil.rmtree(self.temp_dir)
            except:
                pass
        
        # 同时清理过期的图片缓存 (超过1小时的)
        import time
        now = time.time()
        if os.path.exists(self.cache_dir):
            for f in os.listdir(self.cache_dir):
                f_path = os.path.join(self.cache_dir, f)
                try:
                    if os.path.isfile(f_path) and now - os.path.getmtime(f_path) > 3600:
                        os.unlink(f_path)
                except:
                    pass
    
    def download_docx(self, url: str) -> str:
        """从URL下载docx文件到临时目录"""
        try:
            # 尽可能使用 https
            if url.startswith("http://"):
                url = url.replace("http://", "https://", 1)
                
            headers = {
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'
            }
            response = requests.get(url, stream=True, headers=headers, timeout=30)
            response.raise_for_status()
            
            # 从URL获取文件名
            parsed_url = urlparse(url)
            filename = os.path.basename(parsed_url.path)
            if not filename or not filename.endswith('.docx'):
                filename = "downloaded_document.docx"
            
            file_path = os.path.join(self.temp_dir, filename)
            
            with open(file_path, 'wb') as f:
                for chunk in response.iter_content(chunk_size=8192):
                    f.write(chunk)
            
            return file_path
        except Exception as e:
            raise Exception(f"下载文件失败: {str(e)}")
    
    def extract_styles_from_xml(self, docx_path: str):
        """从XML中提取样式信息"""
        try:
            with zipfile.ZipFile(docx_path) as docx_zip:
                # 检查styles.xml是否存在
                if 'word/styles.xml' not in docx_zip.namelist():
                    return
                
                # 读取styles.xml
                styles_xml = docx_zip.read('word/styles.xml')
                root = ET.fromstring(styles_xml)
                
                # 命名空间
                namespaces = {
                    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
                }
                
                # 提取默认字体
                default_font = None
                doc_defaults = root.find('.//w:docDefaults', namespaces)
                if doc_defaults is not None:
                    r_pr_default = doc_defaults.find('.//w:rPrDefault', namespaces)
                    if r_pr_default is not None:
                        r_pr = r_pr_default.find('.//w:rPr', namespaces)
                        if r_pr is not None:
                            font_elem = r_pr.find('.//w:rFonts', namespaces)
                            if font_elem is not None:
                                # 尝试多种字体属性
                                for font_attr in ['ascii', 'eastAsia', 'hAnsi', 'cs']:
                                    font_val = font_elem.get(f'{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{font_attr}')
                                    if font_val:
                                        default_font = font_val
                                        break
                
                # 存储默认字体
                if default_font:
                    self.style_map['_default_font'] = {'font': default_font}
                
                # 遍历所有样式
                for style in root.findall('.//w:style', namespaces):
                    style_id = style.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}styleId')
                    style_type = style.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type')
                    
                    if style_id:
                        # 提取段落样式
                        style_info = {'type': style_type}
                        
                        # 字体 - 改进版，支持多种字体类型
                        font_elem = style.find('.//w:rPr/w:rFonts', namespaces)
                        if font_elem is not None:
                            # 尝试多种字体属性，优先级：eastAsia > ascii > hAnsi > cs
                            font_priority = ['eastAsia', 'ascii', 'hAnsi', 'cs']
                            for font_attr in font_priority:
                                font_val = font_elem.get(f'{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{font_attr}')
                                if font_val:
                                    style_info['font'] = font_val
                                    style_info[f'{font_attr}_font'] = font_val
                                    break
                        
                        # 字号
                        size_elem = style.find('.//w:rPr/w:sz', namespaces)
                        if size_elem is not None:
                            size_val = size_elem.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
                            if size_val:
                                style_info['size'] = float(size_val) / 2  # XML中的字号是实际值的两倍
                        
                        # 加粗
                        bold_elem = style.find('.//w:rPr/w:b', namespaces)
                        if bold_elem is not None:
                            style_info['bold'] = True
                        
                        # 斜体
                        italic_elem = style.find('.//w:rPr/w:i', namespaces)
                        if italic_elem is not None:
                            style_info['italic'] = True
                        
                        # 下划线
                        underline_elem = style.find('.//w:rPr/w:u', namespaces)
                        if underline_elem is not None:
                            style_info['underline'] = True
                        
                        # 上下标
                        vert_align_elem = style.find('.//w:rPr/w:vertAlign', namespaces)
                        if vert_align_elem is not None:
                            val = vert_align_elem.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
                            if val == 'superscript':
                                style_info['superscript'] = True
                            elif val == 'subscript':
                                style_info['subscript'] = True
                        
                        # 突出显示
                        highlight_elem = style.find('.//w:rPr/w:highlight', namespaces)
                        if highlight_elem is not None:
                            val = highlight_elem.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
                            if val:
                                style_info['highlight'] = val
                        
                        # 段落对齐
                        jc_elem = style.find('.//w:pPr/w:jc', namespaces)
                        if jc_elem is not None:
                            jc_val = jc_elem.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
                            alignment_map = {
                                'left': 'left',
                                'center': 'center',
                                'right': 'right',
                                'both': 'justify',
                                'distribute': 'distribute'
                            }
                            if jc_val in alignment_map:
                                style_info['alignment'] = alignment_map[jc_val]
                        
                        # 样式继承关系
                        based_on = style.find('.//w:basedOn', namespaces)
                        if based_on is not None:
                            based_on_val = based_on.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
                            if based_on_val:
                                style_info['based_on'] = based_on_val
                        
                        self.style_map[style_id] = style_info
        except Exception as e:
            print(f"提取样式信息失败: {str(e)}")
    
    def _get_safe_highlight_color(self, run):
        """安全地获取突出显示颜色，避免 WD_COLOR_INDEX 映射错误"""
        try:
            return run.font.highlight_color
        except Exception:
            # 某些文档可能包含 python-docx 不支持的 XML 颜色映射（如 'none'）
            return None

    def get_inherited_style(self, run, paragraph) -> TextStyle:
        """获取继承后的完整样式"""
        # 安全地获取突出显示颜色
        highlight_color = self._get_safe_highlight_color(run)
        
        # 建立缓存键：段落样式ID + run的显式设置
        explicit_props = (
            run.font.name, 
            run.font.size.pt if run.font.size else None,
            run.font.color.rgb if run.font.color else None,
            run.bold,
            run.italic,
            run.underline,
            run.font.subscript,
            run.font.superscript,
            highlight_color
        )
        style_id = paragraph.style.style_id if hasattr(paragraph, 'style') and paragraph.style else "_default_"
        cache_key = (style_id, explicit_props)
        
        if cache_key in self.style_cache:
            return self.style_cache[cache_key]

        # 创建基础样式
        style = TextStyle()
        
        # 1. 首先从run对象获取显式设置的样式
        if run.font.name:
            style.font = run.font.name
        
        if run.font.size:
            style.size = run.font.size.pt
        
        if run.font.color and run.font.color.rgb:
            style.color = f"#{run.font.color.rgb}"
        
        style.bold = run.bold if run.bold is not None else False
        style.italic = run.italic if run.italic is not None else False
        style.underline = run.underline if run.underline is not None else False
        style.subscript = run.font.subscript if run.font.subscript is not None else False
        style.superscript = run.font.superscript if run.font.superscript is not None else False
        
        if highlight_color is not None:
            # 将枚举值转换为颜色名称
            from docx.enum.text import WD_COLOR_INDEX
            color_map = {
                WD_COLOR_INDEX.YELLOW: 'yellow',
                WD_COLOR_INDEX.BRIGHT_GREEN: 'green',
                WD_COLOR_INDEX.TURQUOISE: 'cyan',
                WD_COLOR_INDEX.PINK: 'magenta',
                WD_COLOR_INDEX.BLUE: 'blue',
                WD_COLOR_INDEX.RED: 'red',
                WD_COLOR_INDEX.DARK_BLUE: 'darkBlue',
                WD_COLOR_INDEX.TEAL: 'darkCyan',
                WD_COLOR_INDEX.GREEN: 'darkGreen',
                WD_COLOR_INDEX.VIOLET: 'darkMagenta',
                WD_COLOR_INDEX.DARK_RED: 'darkRed',
                WD_COLOR_INDEX.DARK_YELLOW: 'darkYellow',
                WD_COLOR_INDEX.GRAY_50: 'darkGray',
                WD_COLOR_INDEX.GRAY_25: 'lightGray',
                WD_COLOR_INDEX.BLACK: 'black'
            }
            style.highlight = color_map.get(highlight_color)
        
        # 2. 如果run没有显式设置字体，尝试从段落样式获取
        if not style.font and style_id != "_default_":
            if style_id in self.style_map and 'font' in self.style_map[style_id]:
                style.font = self.style_map[style_id]['font']
        
        # 3. 如果仍然没有字体，尝试从样式继承链中获取
        if not style.font and style_id != "_default_":
            style.font = self._get_font_from_style_chain(style_id)
        
        # 4. 如果仍然没有字体，使用默认字体或兜底为宋体
        if not style.font:
            if '_default_font' in self.style_map:
                style.font = self.style_map['_default_font']['font']
            else:
                style.font = "宋体"
        
        # 5. 如果run没有显式设置字号，尝试从段落样式获取
        if not style.size and style_id != "_default_":
            if style_id in self.style_map and 'size' in self.style_map[style_id]:
                style.size = self.style_map[style_id]['size']
        
        # 6. 如果run没有显式设置粗体，尝试从段落样式获取
        if not style.bold and style_id != "_default_":
            if style_id in self.style_map and 'bold' in self.style_map[style_id]:
                style.bold = self.style_map[style_id]['bold']
        
        # 7. 如果run没有显式设置斜体，尝试从段落样式获取
        if not style.italic and style_id != "_default_":
            if style_id in self.style_map and 'italic' in self.style_map[style_id]:
                style.italic = self.style_map[style_id]['italic']
        
        # 8. 如果run没有显式设置下划线，尝试从段落样式获取
        if not style.underline and style_id != "_default_":
            if style_id in self.style_map and 'underline' in self.style_map[style_id]:
                style.underline = self.style_map[style_id]['underline']

        # 9. 尝试从段落样式获取上下标和突出显示
        if style_id != "_default_":
            if not style.subscript and style_id in self.style_map and 'subscript' in self.style_map[style_id]:
                style.subscript = self.style_map[style_id]['subscript']
            if not style.superscript and style_id in self.style_map and 'superscript' in self.style_map[style_id]:
                style.superscript = self.style_map[style_id]['superscript']
            if not style.highlight and style_id in self.style_map and 'highlight' in self.style_map[style_id]:
                style.highlight = self.style_map[style_id]['highlight']
        
        # 存入缓存
        self.style_cache[cache_key] = style
        return style
    
    def _get_font_from_style_chain(self, style_id: str, visited=None) -> str:
        """从样式继承链中获取字体"""
        if visited is None:
            visited = set()
        
        # 防止循环继承
        if style_id in visited:
            return None
        visited.add(style_id)
        
        # 检查当前样式是否有字体
        if style_id in self.style_map and 'font' in self.style_map[style_id]:
            return self.style_map[style_id]['font']
        
        # 检查是否有基于的样式
        if style_id in self.style_map and 'based_on' in self.style_map[style_id]:
            based_on = self.style_map[style_id]['based_on']
            return self._get_font_from_style_chain(based_on, visited)
        
        return None
    
    def get_inherited_paragraph_style(self, paragraph) -> ParagraphStyle:
        """获取继承后的完整段落样式"""
        style = ParagraphStyle()
        
        # 1. 首先从段落格式获取显式设置的样式
        if paragraph.paragraph_format.alignment:
            alignment_map = {
                WD_ALIGN_PARAGRAPH.LEFT: "left",
                WD_ALIGN_PARAGRAPH.CENTER: "center",
                WD_ALIGN_PARAGRAPH.RIGHT: "right",
                WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
                WD_ALIGN_PARAGRAPH.DISTRIBUTE: "distribute"
            }
            style.alignment = alignment_map.get(paragraph.paragraph_format.alignment, "left")
        
        if paragraph.paragraph_format.line_spacing:
            style.line_spacing = paragraph.paragraph_format.line_spacing
        
        if paragraph.paragraph_format.space_before:
            style.space_before = paragraph.paragraph_format.space_before.pt
        
        if paragraph.paragraph_format.space_after:
            style.space_after = paragraph.paragraph_format.space_after.pt
        
        if paragraph.paragraph_format.left_indent:
            style.left_indent = paragraph.paragraph_format.left_indent.pt
        
        if paragraph.paragraph_format.right_indent:
            style.right_indent = paragraph.paragraph_format.right_indent.pt
        
        if paragraph.paragraph_format.first_line_indent:
            style.first_line_indent = paragraph.paragraph_format.first_line_indent.pt
        
        # 2. 如果没有显式设置对齐方式，尝试从段落样式获取
        if not style.alignment and hasattr(paragraph, 'style') and paragraph.style:
            style_id = paragraph.style.style_id
            if style_id in self.style_map and 'alignment' in self.style_map[style_id]:
                style.alignment = self.style_map[style_id]['alignment']
        
        return style
    
    def extract_table_data(self, table) -> TableData:
        """提取表格数据，支持合并单元格识别"""
        rows_count = len(table.rows)
        cols_count = len(table.columns) if rows_count > 0 else 0
        cells_data = []
        
        # 用于追踪已经处理过的单元格坐标（处理合并单元格时跳过重复坐标）
        visited = set()
        
        for r in range(rows_count):
            row_cells = []
            for c in range(cols_count):
                if (r, c) in visited:
                    # 这个坐标是某个合并单元格的一部分，已经在之前处理过了
                    row_cells.append(None)
                    continue
                
                try:
                    cell = table.cell(r, c)
                    
                    # 识别纵向和横向合并
                    # python-docx 中，合并后的单元格在不同坐标下返回同一个 cell 对象
                    row_span = 1
                    col_span = 1
                    
                    # 探测横向合并 (col_span)
                    while c + col_span < cols_count:
                        try:
                            if table.cell(r, c + col_span) == cell:
                                col_span += 1
                            else:
                                break
                        except:
                            break
                            
                    # 探测纵向合并 (row_span)
                    while r + row_span < rows_count:
                        try:
                            if table.cell(r + row_span, c) == cell:
                                row_span += 1
                            else:
                                break
                        except:
                            break
                    
                    # 标记这个合并单元格覆盖的所有坐标
                    for dr in range(row_span):
                        for dc in range(col_span):
                            visited.add((r + dr, c + dc))
                    
                    # 提取单元格内容（可能包含多个段落，这里简单提取纯文本，或者可以进一步递归解析）
                    row_cells.append({
                        "text": cell.text,
                        "row_span": row_span,
                        "col_span": col_span
                    })
                except Exception:
                    row_cells.append({"text": "", "row_span": 1, "col_span": 1})
                    
            cells_data.append(row_cells)
        
        return TableData(rows=rows_count, columns=cols_count, cells=cells_data)
    
    def extract_image_data(self, docx_path: str) -> List[ImageData]:
        """提取图片数据，包括二进制内容"""
        images = []
        
        try:
            with zipfile.ZipFile(docx_path) as docx_zip:
                # 检查word/media目录是否存在
                media_files = [f for f in docx_zip.namelist() if f.startswith('word/media/')]
                if not media_files:
                    return images
                
                # 获取关系文件以建立 rId -> 文件路径的映射
                rels_xml = docx_zip.read('word/_rels/document.xml.rels')
                rels_root = ET.fromstring(rels_xml)
                rel_map = {}
                namespaces_rel = {'rel': 'http://schemas.openxmlformats.org/package/2006/relationships'}
                for rel in rels_root.findall('.//rel:Relationship', namespaces_rel):
                    rel_map[rel.get('Id')] = rel.get('Target')

                # 获取document.xml以查找图片引用
                document_xml = docx_zip.read('word/document.xml')
                root = ET.fromstring(document_xml)
                
                # 命名空间
                namespaces = {
                    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
                    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
                    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
                    'pic': 'http://schemas.openxmlformats.org/drawingml/2006/picture',
                    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
                }
                
                # 查找所有图片元素
                for drawing in root.findall('.//w:drawing', namespaces):
                    try:
                        # 获取图片尺寸
                        extent = drawing.find('.//wp:extent', namespaces)
                        if extent is not None:
                            width = int(extent.get('cx')) / 914400.0 * 25.4  # EMU to mm
                            height = int(extent.get('cy')) / 914400.0 * 25.4
                        else:
                            width = height = 0
                        
                        # 检查是否是内联图片
                        inline_elem = drawing.find('.//wp:inline', namespaces)
                        inline = inline_elem is not None
                        
                        # 尝试获取对齐方式 (针对非内联/锚定图片)
                        alignment = None
                        align_elem = drawing.find('.//wp:align', namespaces)
                        if align_elem is not None:
                            alignment = align_elem.text # 'center', 'left', 'right' etc.
                        
                        # 尝试获取alt文本
                        doc_pr = drawing.find('.//wp:docPr', namespaces)
                        alt_text = doc_pr.get('descr') if doc_pr is not None else None
                        
                        # 获取图片的 rId
                        blip = drawing.find('.//a:blip', namespaces)
                        binary_content = None
                        image_path = None
                        cached_path = None
                        
                        if blip is not None:
                            rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                            if rid in rel_map:
                                target = rel_map[rid]
                                # 处理相对路径，通常是 'media/image1.png'
                                if target.startswith('media/'):
                                    full_path = f"word/{target}"
                                else:
                                    full_path = target
                                
                                if full_path in docx_zip.namelist():
                                    img_data = docx_zip.read(full_path)
                                    binary_content = base64.b64encode(img_data).decode('utf-8')
                                    image_path = full_path
                                    
                                    # 缓存到本地文件
                                    ext = os.path.splitext(full_path)[1] or '.png'
                                    cache_filename = f"img_{rid}_{len(images)}{ext}"
                                    cached_path = os.path.join(self.cache_dir, cache_filename)
                                    try:
                                        with open(cached_path, 'wb') as img_f:
                                            img_f.write(img_data)
                                    except Exception as e:
                                        print(f"缓存图片失败: {str(e)}")
                                        cached_path = None

                        images.append(ImageData(
                            width=width,
                            height=height,
                            inline=inline,
                            alignment=alignment,
                            alt_text=alt_text,
                            image_path=cached_path if cached_path else image_path,
                            binary_content=binary_content
                        ))
                    except Exception:
                        continue
        except Exception:
            pass
        
        return images
    
    def parse(self, local_path: str) -> Dict[str, Any]:
        """解析本地docx文件并返回结构化数据"""
        try:
            # 打开文档
            doc = Document(local_path)
            
            # 建立映射索引，消除嵌套循环 (O(N^2) -> O(N))
            para_map = {p._element: p for p in doc.paragraphs}
            table_map = {t._element: t for t in doc.tables}
            
            # 提取图片数据
            images = self.extract_image_data(local_path)
            
            # 初始化结果
            result = {
                'document_elements': [],
                'images': [],
                'metadata': {
                    'paragraphs_count': 0,
                    'tables_count': 0,
                    'images_count': len(images)
                },
                'styles': self.style_map  # 添加样式信息到结果中
            }
            
            # 处理文档元素
            image_index = 0
            for element in doc.element.body:
                if element.tag.endswith('p'):  # 段落
                    paragraph = para_map.get(element)
                    
                    if paragraph:
                        # 提前提取段落样式，用于可能的图片对齐
                        para_style = self.get_inherited_paragraph_style(paragraph)

                        # 检查段落中是否包含图片
                        drawing = element.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing')
                        if drawing is not None and image_index < len(images):
                            img_data = images[image_index]
                            # 优先使用图片自身的对齐方式，如果没有，则继承段落的对齐方式
                            if not img_data.alignment and para_style and para_style.alignment:
                                img_data.alignment = para_style.alignment
                            
                            result['document_elements'].append({
                                'element_type': 'image',
                                'content': img_data.__dict__
                            })
                            image_index += 1
                            continue # 如果是纯图片段落，处理完就跳过
                        
                        # 提取文本和字符样式（使用增强版方法）
                        text_content = []
                        for run in paragraph.runs:
                            if run.text:
                                text_style = self.get_inherited_style(run, paragraph)
                                text_content.append({
                                    'text': run.text,
                                    'style': text_style.__dict__
                                })
                        
                        # 无论是否包含文本内容，都保留该段落（支持空行识别）
                        result['document_elements'].append({
                            'element_type': 'paragraph',
                            'content': text_content,
                            'paragraph_style': para_style.__dict__
                        })
                        result['metadata']['paragraphs_count'] += 1
                
                elif element.tag.endswith('tbl'):  # 表格
                    table = table_map.get(element)
                    
                    if table:
                        table_data = self.extract_table_data(table)
                        result['document_elements'].append({
                            'element_type': 'table',
                            'content': table_data.__dict__
                        })
                        result['metadata']['tables_count'] += 1
                
                elif element.tag.endswith('r'):  # 可能包含图片
                    # 检查是否是图片引用
                    drawing = element.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing')
                    if drawing is not None and image_index < len(images):
                        result['document_elements'].append({
                            'element_type': 'image',
                            'content': images[image_index].__dict__
                        })
                        image_index += 1
            
            # 添加图片数据
            result['images'] = [img.__dict__ for img in images]
            
            return result
        except Exception as e:
            raise Exception(f"解析文件内容失败: {str(e)}")

    def parse_docx(self, url: str) -> Dict[str, Any]:
        """解析docx文件并返回结构化数据"""
        # 下载文件
        docx_path = self.download_docx(url)
        
        try:
            # 提取样式信息
            self.extract_styles_from_xml(docx_path)
            
            # 使用本地解析逻辑
            return self.parse(docx_path)
        
        finally:
            # 清理临时文件
            try:
                if os.path.exists(docx_path):
                    os.unlink(docx_path)
                if os.path.exists(self.temp_dir) and not os.listdir(self.temp_dir):
                    os.rmdir(self.temp_dir)
            except:
                pass


def main():
    """主函数，用于测试"""
    import json
    
    if len(os.sys.argv) < 2:
        print("用法: python enhanced_docx_parser.py <docx_url>")
        return
    
    url = os.sys.argv[1]
    parser = EnhancedDocxParser()
    
    try:
        result = parser.parse_docx(url)
        print(json.dumps(result, indent=2, ensure_ascii=False))
    except Exception as e:
        print(f"解析失败: {str(e)}")


if __name__ == "__main__":
    main()